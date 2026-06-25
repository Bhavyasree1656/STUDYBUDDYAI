"""Local text extraction — no API calls."""
import io
from pathlib import Path


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    ext = Path(filename).suffix.lower().lstrip(".")
    extractors = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "txt": _extract_txt,
        "pptx": _extract_pptx,
    }
    if ext not in extractors:
        raise ValueError(f"Unsupported file type: .{ext}")
    return extractors[ext](file_bytes)


def _extract_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=data, filetype="pdf")
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        if slide_texts:
            parts.append(f"[Slide {i}]\n" + "\n".join(slide_texts))
    return "\n\n".join(parts)
